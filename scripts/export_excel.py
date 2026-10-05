"""
รวมข้อมูลที่ใช้ทั้งหมดเป็นไฟล์ Excel ไฟล์เดียวสำหรับส่ง → data/ข้อมูลที่ใช้_ถังรวม.xlsx
รัน:  python scripts/export_excel.py   (ต้องมี openpyxl)
"""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "ข้อมูลที่ใช้_ถังรวม.xlsx"
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3B57")
NUMERIC = {"size_kg", "cylinders_at_start", "full_out", "empty_in", "gas_amount", "deposit_amount",
           "paid_amount", "amount", "gas_price", "cylinder_value", "share_in_shop", "market_share_pct"}

SHEETS = [
    ("customers", "data/raw/customers.csv", "ลูกค้า 220 ราย (ข้อมูลจำลอง ชื่อสมมติ)"),
    ("transactions", "data/raw/transactions.csv", "รายการขาย แลกถัง ซื้อถังใหม่ คืนถัง (ข้อมูลจำลอง)"),
    ("payments", "data/raw/payments.csv", "การชำระเงินของลูกค้าเครดิต (ข้อมูลจำลอง)"),
    ("products", "data/raw/products.csv", "ราคาแก๊สและมูลค่าถัง (ค่าประมาณ)"),
    ("brands", "data/raw/brands.csv", "แบรนด์ที่ร้านขาย และสัดส่วนในร้าน (สมมติฐาน)"),
    ("market_share", "data/reference/market_share_lpg_q2_2568.csv", "ส่วนแบ่งตลาด LPG ไตรมาส 2/2568 (ข้อมูลจริง)"),
]


def read(p):
    with open(ROOT / p, encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HEAD
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    info = wb.active
    info.title = "อ่านก่อน"
    sizes = {}
    for name, path, _ in SHEETS:
        rows = read(path)
        ws = wb.create_sheet(name)
        header = rows[0]
        ws.append(header)
        for r in rows[1:]:
            ws.append([(float(v) if "." in v else int(v)) if h in NUMERIC and v not in ("",) else v
                       for h, v in zip(header, r)])
        style_header(ws, len(header))
        for i, h in enumerate(header, 1):
            width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows[1:200])) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name=F)
        sizes[name] = len(rows) - 1

    # ---- สรุปตัวเลข (สูตร คำนวณจากชีตข้อมูล)
    s = wb.create_sheet("สรุปตัวเลข", 1)
    n = sizes["transactions"] + 1
    T = lambda col: f"transactions!${col}$2:${col}${n}"
    s.append(["ตัวชี้วัด", "ค่า", "วิธีคำนวณ"])
    s.append(["จำนวนลูกค้า", f"=COUNTA(customers!A2:A{sizes['customers'] + 1})", "นับรหัสลูกค้า"])
    s.append(["จำนวนรายการ", f"=COUNTA({T('A')})", "นับรายการใน transactions"])
    s.append(["จำนวนออเดอร์ (มีส่งถังเต็ม)", f'=COUNTIF({T("H")},">0")', "full_out > 0"])
    s.append(["ยอดขายค่าแก๊ส (บาท)", f"=SUM({T('J')})", "รวม gas_amount"])
    s.append(["ถังที่อยู่กับลูกค้า (ใบ)",
              f"=SUM(customers!F2:F{sizes['customers'] + 1})+SUM({T('H')})-SUM({T('I')})",
              "ถังตอนเริ่ม + ถังเต็มส่งออก − ถังเปล่ารับคืน"])
    s.append(["ยอดขายเชื่อ (บาท)", f'=SUMIFS({T("J")},{T("L")},"credit")', "gas_amount ที่ payment_method = credit"])
    s.append(["รับชำระจากลูกค้าเครดิต (บาท)", f"=SUM(payments!D2:D{sizes['payments'] + 1})", "รวม payments"])
    s.append(["ลูกหนี้คงค้าง (บาท)", "=B7-B8", "ขายเชื่อ − รับชำระ"])
    s.append([])
    s.append(["ยอดขายค่าแก๊สแยกแบรนด์", "บาท", "สัดส่วน"])
    first = s.max_row + 1
    for i in range(sizes["brands"]):
        r = s.max_row + 1
        s.append([f"=brands!B{i + 2}", f'=SUMIFS({T("J")},{T("E")},brands!A{i + 2})', f"=IF($B$5=0,0,B{r}/$B$5)"])
        s.cell(row=r, column=3).number_format = "0.0%"
    for r in range(2, s.max_row + 1):
        s.cell(row=r, column=2).number_format = "#,##0" if s.cell(row=r, column=3).number_format != "0.0%" else "#,##0"
    for row in s.iter_rows():
        for cell in row:
            cell.font = Font(name=F, bold=cell.row in (1, first - 1))
    for c in range(1, 4):
        s.cell(row=1, column=c).font = Font(name=F, bold=True, color="FFFFFF")
        s.cell(row=1, column=c).fill = HEAD
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 34, 16, 44

    # ---- อ่านก่อน
    lines = [
        ["ข้อมูลที่ใช้ทำ Dashboard ร้านแก๊ส LPG ที่ขายหลายแบรนด์"],
        ["วิชา Business Idea Creation | บริษัท ถังรวม"],
        ["ช่วงข้อมูล 1 เม.ย. 2569 - 30 ก.ย. 2569"],
        [],
        ["ที่มาของข้อมูล"],
        ["ค้นหาใน Kaggle และแหล่งข้อมูลเปิดแล้ว ไม่พบข้อมูลรายการขายของร้านแก๊ส"],
        ["จึงสร้างข้อมูลจำลองด้วย scripts/generate_data.py (รันซ้ำได้ผลเหมือนเดิม)"],
        ["ข้อมูลจริงมีชีตเดียวคือ market_share: กรมธุรกิจพลังงาน อ้างในรายงาน LH Bank ต.ค. 2568"],
        ["https://www.lhbank.co.th/getattachment/031004c9-34dc-4c54-b374-4928bb44246c/economic-analysis-Industry-Outlook-2025-Natural-Gas_Oct_2025"],
        ["แถว OTHERS ในชีต market_share คำนวณจาก 100 ลบผลรวม 5 รายแรก"],
        [],
        ["ชีต", "จำนวนแถว", "คำอธิบาย"],
        ["สรุปตัวเลข", "", "ตัวเลขสำคัญ คำนวณด้วยสูตรจากชีตข้อมูล"],
    ] + [[name, sizes[name], desc] for name, _, desc in SHEETS] + [
        [],
        ["ความหมายคอลัมน์ที่สำคัญ (ดูทั้งหมดใน data/DATA_DICTIONARY.md)"],
        ["brand", "", "WP = เวิลด์แก๊ส, PTT = ปตท., UNQ = ยูนิคแก๊ส, SGP = สยามแก๊ส"],
        ["txn_type", "", "exchange = แลกถัง, new_cylinder = ซื้อถังใหม่ (มัดจำ), return_cylinder = คืนถัง"],
        ["full_out / empty_in", "", "ถังเต็มที่ส่งออก / ถังเปล่าที่ได้คืน"],
        ["payment_method", "", "cash = เงินสด, transfer = โอน, credit = ขายเชื่อ (ยังไม่จ่าย)"],
        ["cylinders_at_start", "", "ถังที่ลูกค้าถืออยู่ ณ 1 เม.ย. 2569"],
        [],
        ["ข้อจำกัด: ข้อมูลร้านเป็นข้อมูลจำลอง ราคาแก๊สและมูลค่าถังเป็นค่าประมาณ ไม่ใช่ราคาประกาศ"],
    ]
    for l in lines:
        info.append(l)
    for row in info.iter_rows():
        for cell in row:
            cell.font = Font(name=F)
    info["A1"].font = Font(name=F, bold=True, size=14)
    for r in (5, 12):
        for c in range(1, 4):
            info.cell(row=r, column=c).font = Font(name=F, bold=True)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width, info.column_dimensions["C"].width = 24, 12, 70

    wb.save(OUT)
    print(OUT.relative_to(ROOT), sizes)


if __name__ == "__main__":
    main()
